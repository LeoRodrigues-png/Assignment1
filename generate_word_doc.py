from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from pathlib import Path


def build_document(out_path: Path) -> None:
    doc = Document()

    # Title
    p = doc.add_paragraph()
    run = p.add_run("Web Services - Assignment 1\nInventory Management API")
    run.bold = True
    run.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p2 = doc.add_paragraph()
    run2 = p2.add_run("(FastAPI + MongoDB + Docker + Jenkins + Postman/Newman + Monitoring)")
    run2.italic = True
    run2.font.size = Pt(12)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    # Student details placeholders
    for label in ["Student Name", "Student ID", "Date"]:
        pp = doc.add_paragraph()
        pp.add_run(f"{label}: ").bold = True

    def add_heading(text: str) -> None:
        doc.add_heading(text, level=1)

    def add_subheading(text: str) -> None:
        doc.add_heading(text, level=2)

    def add_bullets(lines: list[str]) -> None:
        for line in lines:
            doc.add_paragraph(line, style="List Bullet")

    # 1 Intro
    add_heading("1. Introduction / Overview")
    doc.add_paragraph(
        "This project provides an inventory management API. Product data is stored in MongoDB and exposed through a FastAPI service. "
        "The workflow uses Docker for repeatability, Jenkins for automated build/test, Postman/Newman for endpoint testing, "
        "and Prometheus/Grafana for monitoring (/metrics is exposed by the API)."
    )

    add_subheading("Tech stack used")
    add_bullets(
        [
            "FastAPI (API + automatic validation + docs at /docs)",
            "Pydantic models for request/response validation",
            "MongoDB for persistent inventory storage",
            "Docker for containerised deployment",
            "Jenkins pipeline to build/run/test + generate README + zip artifact",
            "Postman collection exported to JSON + Newman to run in Jenkins",
            "Prometheus + Grafana for monitoring (extra credit)",
        ]
    )

    # 2 Structure
    add_heading("2. Project Files / Structure")
    doc.add_paragraph("Main folder layout:")
    doc.add_paragraph(
        "assignment1/\n"
        "- app/\n"
        "- scripts/\n"
        "- tests/\n"
        "- monitoring/\n"
        "- Dockerfile\n"
        "- docker-compose.yml\n"
        "- Jenkinsfile\n"
        "- requirements.txt\n"
        "- products.csv\n"
        "- generate_readme.py\n"
        "- README.txt (generated)\n"
    )
    doc.add_paragraph("Screenshot needed (optional): a brief file explorer view of the project root.")

    # 3 Data setup
    add_heading("3. Data Setup (CSV -> MongoDB)")
    doc.add_paragraph(
        "The script reads `products.csv`, converts each row to a Python dict, and inserts/upserts the documents into MongoDB. "
        "Re-running the script is safe because it uses `upsert` by `product_id`."
    )
    doc.add_paragraph("Key file:")
    doc.add_paragraph("- `scripts/import_csv_to_mongo.py`")

    doc.add_paragraph("Screenshots to include (place in this order):")
    add_bullets(
        [
            "Docker containers running (show Mongo + API): run `docker compose up -d --build` and screenshot `docker compose ps`.",
            "Optional: screenshot the API logs showing import completion (from `docker compose logs api --tail=20`).",
        ]
    )

    # 4 FastAPI
    add_heading("4. FastAPI API Implementation")
    doc.add_paragraph(
        "FastAPI serves endpoints defined below. Pydantic models and typed query parameters enforce validation for inputs. "
        "Swagger UI is available at `/docs` and helps demonstrate the endpoints interactively."
    )

    doc.add_paragraph("Key files:")
    add_bullets(
        [
            "`app/main.py` (endpoints + metrics)",
            "`app/models.py` (Pydantic schemas)",
            "`app/database.py` (Mongo connection)",
            "`app/utils.py` (USD->EUR conversion)",
        ]
    )

    add_subheading("Screenshots of all endpoints running")
    doc.add_paragraph("Include screenshots exactly in this sequence:")

    endpoint_placeholders = [
        "1) Swagger UI overview at http://localhost:8000/docs (show all endpoints listed).",
        "2) `GET /getAll` response (show JSON list).",
        "3) `GET /getSingleProduct?product_id=1001` response (existing product).",
        "4) `GET /startsWith?letter=S` response (names starting with S, case-insensitive).",
        "5) `GET /paginate?start_id=1001&end_id=1100` response (must show <= 10 items in the range).",
        "6) `GET /convert?product_id=1001` response (show `unit_price_eur` and the rate).",
        "7) `POST /addNew` success response (HTTP 201) with a demo product_id not in the CSV (e.g., 9999).",
        "8) `DELETE /deleteOne?product_id=9999` success response (HTTP 200 deleted=true).",
        "9) Validation proof: show one 422 Validation Error from Swagger (e.g., pass a non-integer to `product_id`).",
    ]
    for item in endpoint_placeholders:
        doc.add_paragraph(item, style="List Number")

    # 5 Testing
    add_heading("5. Testing (Postman + Newman)")
    doc.add_paragraph(
        "A Postman collection is exported to JSON and executed using Newman inside the Jenkins pipeline. "
        "This demonstrates that all endpoints work end-to-end against the running Docker containers."
    )
    doc.add_paragraph("Key file:")
    doc.add_paragraph("- `tests/postman_collection.json`")

    add_subheading("Screenshots to include")
    add_bullets(
        [
            "Screenshot of Newman run output in the Jenkins console (or terminal).",
            "Screenshot showing the Postman collection contains all requests (optional if your marking scheme is strict).",
        ]
    )

    # 6 Docker
    add_heading("6. Docker Configuration")
    doc.add_paragraph("Docker is used for repeatable deployment and to match the DevOps requirement.")
    add_bullets(["`Dockerfile`", "`docker-compose.yml` (runs Mongo + API, imports CSV on startup)"])

    add_subheading("Screenshots to include")
    add_bullets(
        [
            "`docker compose up -d --build` terminal output (containers started).",
            "`docker compose ps` output (API and Mongo status).",
        ]
    )

    # 7 Jenkins
    add_heading("7. Jenkins Pipeline")
    doc.add_paragraph(
        "The Jenkins pipeline builds the API Docker image, starts MongoDB, runs the API container in the background, "
        "executes Newman tests, generates README.txt, and finally creates a zip artifact named like `complete-DATE-TIME.zip`."
    )

    doc.add_paragraph("Key file:")
    doc.add_paragraph("- `Jenkinsfile`")

    add_subheading("Screenshots to include")
    add_bullets(
        [
            "Jenkins pipeline stage view (show green stages).",
            "Console output showing Newman run succeeded.",
            "Console output showing README.txt generated.",
            "Console output showing zip artifact created.",
        ]
    )

    # 8 Monitoring
    add_heading("8. Monitoring / Management (Extra Credit)")
    doc.add_paragraph(
        "Prometheus scrapes the API's metrics endpoint (`/metrics`) and Grafana visualises them. "
        "This helps verify the API stays online and provides basic performance monitoring."
    )

    add_subheading("Key files")
    add_bullets(["`monitoring/prometheus.yml`", "`monitoring/docker-compose.monitoring.yml`"])

    add_subheading("Screenshots to include")
    add_bullets(
        [
            "API `/metrics` output screenshot (http://localhost:8000/metrics).",
            "Prometheus Targets page screenshot showing the API target is UP (http://localhost:9090).",
            "Grafana login/datasource screenshot (http://localhost:3000).",
            "Grafana dashboard/panel screenshot showing a metrics graph/query.",
        ]
    )

    # 9 Zip output
    add_heading("9. Final Zip Output")
    doc.add_paragraph(
        "After Jenkins finishes, it creates a zip file containing the API source code and Docker configuration. "
        "Screenshot the zip contents as requested by the brief."
    )

    add_subheading("Screenshots to include")
    add_bullets(
        [
            "Screenshot of the Jenkins artifacts showing `complete-DATE-TIME.zip`.",
            "Screenshot of the zip file contents listing folders/files.",
        ]
    )

    # 10 references
    add_heading("10. Appendix / References")
    doc.add_paragraph("FastAPI documentation: https://fastapi.tiangolo.com/")
    doc.add_paragraph("MongoDB docs: https://www.mongodb.com/docs/")
    doc.add_paragraph("Docker docs: https://docs.docker.com/")
    doc.add_paragraph("Prometheus docs: https://prometheus.io/docs/")
    doc.add_paragraph("Grafana docs: https://grafana.com/docs/")

    doc.save(out_path)


def main() -> None:
    root = Path(r"C:\\Users\\leona\\OneDrive\\Uni\\Year 4\\Semester 2\\Web Services\\Assignment1")
    out_path = root / "Assignment1_Inventory_API_Document.docx"
    build_document(out_path)
    print(f"Created: {out_path}")


if __name__ == "__main__":
    main()

